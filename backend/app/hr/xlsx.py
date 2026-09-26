"""Import/export of the enterprise HR workbook shape."""

from datetime import date, datetime
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook


HR_HEADERS = (
    "ФИО", "Должность", "Должность англ.", "Отдел", "Руководитель", "Юр.лицо",
    "Формат работы", "Дата выхода", "Дата окончания ИС", "Пол", "ДР Месяц",
    "День рождения", "Визитка", "ТГ", "Телефон", "Почта", "Образование",
    "Уч. степень", "Специальность", "Рекомендация", "Рекрутер (HR)",
    "Ссылка на фото", "Ссылка на картинку в почте", "Комментарии", "Страховка",
)

FIELD_BY_INDEX = (
    "name", "position", "position_en", "department", "department_head", "legal_entity",
    "work_format", "hire_date", "probation_end_date", "gender", "birth_month_label",
    "birth_date", "business_card", "telegram", "personal_phone", "work_email",
    "education_institution", "academic_degree", "education_specialty", "recommendation",
    "recruiter", "photo_source_url", "mail_image_url", "comments", "insurance",
)

MAX_XLSX_BYTES = 20 * 1024 * 1024
MAX_UNPACKED_BYTES = 80 * 1024 * 1024


def _iso(value):
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if value is None or value == "":
        return None
    return str(value).strip() or None


def _normalized_header(value) -> str:
    return " ".join(str(value or "").split()).casefold()


def parse_hr_xlsx(blob: bytes) -> list[dict]:
    if not blob or len(blob) > MAX_XLSX_BYTES:
        raise ValueError("Кадровый XLSX пуст или превышает 20 МБ")
    try:
        with ZipFile(BytesIO(blob)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_UNPACKED_BYTES:
                raise ValueError("Распакованный XLSX превышает допустимый размер")
    except BadZipFile as error:
        raise ValueError("Файл не является XLSX") from error
    try:
        workbook = load_workbook(BytesIO(blob), read_only=True, data_only=True)
    except Exception as error:
        raise ValueError("Не удалось прочитать кадровый XLSX") from error
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    header = next(iterator, ())
    actual = tuple(_normalized_header(value) for value in header[:25])
    expected = tuple(_normalized_header(value) for value in HR_HEADERS)
    if actual != expected:
        raise ValueError("Первая строка XLSX не соответствует 25 колонкам кадрового шаблона")
    rows = []
    for row_number, values in enumerate(iterator, 2):
        cells = list(values[:25]) + [None] * max(0, 25 - len(values))
        if not any(value not in (None, "") for value in cells):
            continue
        item = {field: _iso(cells[index]) for index, field in enumerate(FIELD_BY_INDEX)}
        item["row_number"] = row_number
        rows.append(item)
    workbook.close()
    return rows


def export_hr_xlsx(employees: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Кадровый учет"
    sheet.append(HR_HEADERS)
    for employee in employees:
        values = {
            "name": employee.get("plan_name"), "position": employee.get("position"),
            "position_en": employee.get("position_en"), "department": employee.get("department"),
            "department_head": employee.get("department_head"), "legal_entity": employee.get("legal_entity"),
            "work_format": employee.get("work_format"), "hire_date": employee.get("hire_date"),
            "probation_end_date": employee.get("probation_end_date"), "gender": employee.get("gender"),
            "birth_month_label": employee.get("birth_month_label"), "birth_date": employee.get("birth_date"),
            "business_card": employee.get("business_card"), "telegram": employee.get("telegram"),
            "personal_phone": employee.get("personal_phone"), "work_email": employee.get("work_email"),
            "education_institution": employee.get("education_institution"), "academic_degree": employee.get("academic_degree"),
            "education_specialty": employee.get("education_specialty"), "recommendation": employee.get("recommendation"),
            "recruiter": employee.get("recruiter"), "photo_source_url": employee.get("photo_source_url"),
            "mail_image_url": employee.get("mail_image_url"), "comments": employee.get("comments"),
            "insurance": employee.get("insurance"),
        }
        sheet.append([values[field] for field in FIELD_BY_INDEX])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:Y{max(1, sheet.max_row)}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
