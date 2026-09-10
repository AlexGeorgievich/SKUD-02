from openpyxl import Workbook
from ...domain.constants import CODES
from ...domain.calendar import month_dates
from .formatting import style, bytes_wb

def export(result,mapping=None,department='',employee='',view='all'):
    wb=Workbook();wb.remove(wb.active)
    em=[e for e in result['employees'] if (not department or e['department']==department) and (not employee or e['id']==employee)]
    ids={e['id'] for e in em};days=[d for d in result['days'] if d['id'] in ids]
    name=lambda i:(mapping or {}).get(i,i)
    tables={
      'За месяц':[['Отдел','Сотрудник']+CODES+['Дни с регистрацией','Часы присутствия СКУД','Дни с замечаниями']]+[[e['department'],name(e['id'])]+[e['counts'].get(c,0) for c in CODES]+[e['registered'],round(e['minutes']/60,2),e['issues']] for e in em],
      'По дням':[['Отдел','Сотрудник','Дата','План','Приход','Уход','Часы присутствия СКУД','Результат','Исходная ячейка']]+[[d['department'],name(d['id']),d['date'],d['plan'],d['arrival'],d['departure'],round(d['minutes']/60,2) if d['minutes'] is not None else None,d['result'],d['raw']] for d in days],
      'Проверка данных':[['Отдел','Сотрудник','Дата','Замечание']]+[[d['department'],name(d['id']),d['date'],d['result']] for d in days if d['problem']],
    }
    dates=month_dates(result['period']);grid=[['Отдел','Сотрудник','Строка']+[d.day for d in dates]]
    for e in em:
        own=[d for d in days if d['id']==e['id']]
        grid.append([e['department'],name(e['id']),'План']+[d['plan'] for d in own])
        grid.append([e['department'],name(e['id']),'Факт']+[d['raw'] or '—' for d in own])
    tables['План-факт']=grid
    for title in ['План-факт','За месяц','Проверка данных','По дням']:
        if view!='all' and view!=title:continue
        ws=wb.create_sheet(title)
        for row in tables[title]:
            ws.append(row)
            for c in ws[ws.max_row]:
                if isinstance(c.value,str):c.data_type='s' # prevents formula injection
        style(ws,1,title=='План-факт')
        if title=='По дням':ws.column_dimensions['H'].width=65;ws.column_dimensions['I'].width=22
        if title=='Проверка данных':ws.column_dimensions['D'].width=80
    if not wb.worksheets:raise ValueError('Неизвестный отчёт')
    return bytes_wb(wb)

