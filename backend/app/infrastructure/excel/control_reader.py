import io,zipfile
from openpyxl import load_workbook

CONTROL_HEADERS=('Отдел','Сотрудник','Активное рабочее время (часов)','Полезная активность (часов)','Прогулы/переработки','Отпуск','Замечания (нарушения, подозрительная активность)','Посещения hh.ru (минуты)')

def _header(value):return ' '.join(str(value or '').split()).casefold()
def read_control_xlsx(blob:bytes,period:str)->list[dict]:
 if len(blob)>20*1024*1024:raise ValueError('Файл превышает 20 МБ')
 try:
  with zipfile.ZipFile(io.BytesIO(blob)) as z:
   if sum(x.file_size for x in z.infolist())>80*1024*1024:raise ValueError('Слишком большой распакованный Excel')
  workbook=load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
 except Exception as error:raise ValueError('Не удалось прочитать XLSX') from error
 sheet=workbook.active;iterator=sheet.iter_rows(values_only=True);headers=next(iterator,())
 expected=[_header(x) for x in CONTROL_HEADERS]
 actual=[_header(x).split(' за ')[0] for x in headers[:8]]
 if actual[0:2]!=expected[0:2] or actual[2]!=expected[2] or actual[3:]!=expected[3:]:raise ValueError('Неподдерживаемые заголовки отчёта СК')
 result=[];department=''
 for row_number,row in enumerate(iterator,2):
  cells=list(row[:8])+[None]*max(0,8-len(row));errors=[]
  if cells[0] not in (None,''):department=str(cells[0]).strip()
  name=str(cells[1] or '').strip()
  if not name:errors.append({'field':'name','error':'Не указан сотрудник'})
  numbers=[]
  for index,field in ((2,'active_hours'),(3,'useful_hours'),(7,'hh_minutes')):
   value=cells[index]
   try:numbers.append(float(value) if value not in (None,'') else None)
   except (TypeError,ValueError):numbers.append(None);errors.append({'field':field,'error':'Ожидается число'})
  result.append({'row_number':row_number,'department':department,'name':name,'active_hours':numbers[0],'useful_hours':numbers[1],'attendance_note':cells[4],'vacation':cells[5],'remarks':cells[6],'hh_minutes':numbers[2],'errors':errors})
 workbook.close();return result
