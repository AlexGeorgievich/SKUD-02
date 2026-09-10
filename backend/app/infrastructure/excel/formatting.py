import io
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from ...domain.constants import COLORS

def style(ws,header=1,wide=False):
    ws.sheet_view.showGridLines=False
    ws.freeze_panes=f'C{header+1}'
    ws.auto_filter.ref=f'A{header}:{get_column_letter(ws.max_column)}{ws.max_row}'
    ws.print_title_rows=f'1:{header}'
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape'
    ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1
    ws.page_setup.fitToHeight=0
    ws.oddFooter.center.text='PPL Group | TimeTrack Pro | &P / &N'
    for row in ws:
        for c in row:
            c.font=Font(name='Arial',size=10,color='15213B')
            c.alignment=Alignment(vertical='center',wrap_text=True)
            if c.row==header:
                c.fill=PatternFill('solid',fgColor='1E3A5F')
                c.font=Font(name='Arial',size=10,bold=True,color='FFFFFF')
            elif str(c.value) in COLORS:
                c.fill=PatternFill('solid',fgColor=COLORS[c.value])
            if isinstance(c.value,(float,int)): c.number_format='0.00' if isinstance(c.value,float) else '0'
    for col in range(1,ws.max_column+1):
        ws.column_dimensions[get_column_letter(col)].width=24 if col<=2 else (17 if wide else 23)
    ws.row_dimensions[header].height=36
    for r in range(header+1,ws.max_row+1): ws.row_dimensions[r].height=52 if wide else 32


def bytes_wb(wb):
    out=io.BytesIO();wb.save(out);return out.getvalue()

