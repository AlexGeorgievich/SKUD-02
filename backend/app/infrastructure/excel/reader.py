import io
from collections import Counter
from openpyxl import load_workbook
from ...domain.calendar import month_dates
from ...domain.identity import normalize

def read_input(blob,kind,period):
    import zipfile
    if len(blob)>20*1024*1024:raise ValueError('Файл превышает 20 МБ')
    try:
        with zipfile.ZipFile(io.BytesIO(blob)) as z:
            if sum(x.file_size for x in z.infolist())>80*1024*1024:raise ValueError('Слишком большой распакованный Excel')
        wb=load_workbook(io.BytesIO(blob),data_only=False,read_only=True)
    except Exception as e:raise ValueError('Не удалось прочитать XLSX: '+str(e)) from e
    ws=wb.active
    if ws.max_row>3000 or ws.max_column>80:raise ValueError('Превышены ограничения: 3000 строк / 80 колонок')
    dates=month_dates(period);header=4 if kind=='plan' else 9
    if ws.cell(2,2).value!=period:raise ValueError(f'{kind}: период в B2 не совпадает с выбранным')
    expected=['Отдел','Сотрудник'] if kind=='plan' else ['№','Ф И']
    if [ws.cell(header,c).value for c in (1,2)]!=expected:raise ValueError(f'{kind}: неподдерживаемый шаблон заголовка (строка {header})')
    for j,d in enumerate(dates,3):
        if str(ws.cell(header,j).value).split('\n')[0]!=str(d.day):raise ValueError('Отсутствуют или переставлены дни месяца')
    rows=[];dept=''
    for row in ws.iter_rows(min_row=header+1,values_only=True):
        if row[0]=='Итого':continue
        if not row[1]:
            if any(v is not None and v!='' for v in row):raise ValueError('Строка данных без ФИ')
            continue
        if row[0] is not None:dept=str(row[0])
        if kind=='plan' and not dept:raise ValueError('Не задан отдел')
        original=str(row[1]);key=normalize(original)
        rows.append(dict(name=original.strip(),normalized=key,department=dept if kind=='plan' else '',values=list(row[2:2+len(dates)])))
    if not rows:raise ValueError('Нет сотрудников')
    duplicates=[k for k,n in Counter(r['normalized'] for r in rows).items() if n>1]
    if duplicates:raise ValueError(f'{kind}: неоднозначные ФИ ({len(duplicates)}). Сверка остановлена; уточните исходные записи.')
    wb.close();return rows

