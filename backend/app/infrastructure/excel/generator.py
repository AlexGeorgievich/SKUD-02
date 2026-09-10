from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from ...domain.constants import DEPTS, CODES
from ...domain.calendar import month_dates
from ...domain.identity import normalize
from .formatting import style, bytes_wb

def generate(period='2026-08',seed=42):
    from faker import Faker
    import random
    fake=Faker('ru_RU');fake.seed_instance(seed);rng=random.Random(seed)
    dates=month_dates(period); used=set();people=[]
    plan=Workbook();p=plan.active;p.title='План'
    p.append(['План работ — PPL Group (синтетические данные)']);p.append(['Период',period]);p.append(['График 5/2: суббота и воскресенье — выходные; без праздничных переносов'])
    p.append(['Отдел','Сотрудник']+[f'{d.day}\n'+['пн','вт','ср','чт','пт','сб','вс'][d.weekday()] for d in dates])
    for dept,count in DEPTS.items():
        first=p.max_row+1
        for _ in range(count):
            while True:
                female=rng.choice([True,False])
                name=(fake.last_name_female()+' '+fake.first_name_female()) if female else (fake.last_name_male()+' '+fake.first_name_male())
                if normalize(name) not in used:used.add(normalize(name));break
            codes=[('Вых' if d.weekday()>4 else rng.choices(CODES[:5],[65,25,5,3,2])[0]) for d in dates]
            people.append((name,codes));p.append([dept,name]+codes)
        if count>1:p.merge_cells(start_row=first,start_column=1,end_row=p.max_row,end_column=1)
    style(p,4,True)
    p.column_dimensions['A'].width=10
    p.column_dimensions['B'].width=32
    for row in p:
        if row[0].row>4 and row[0].value:
            row[0].alignment=Alignment(textRotation=90,vertical='center',horizontal='center')
            if DEPTS[row[0].value]==1:p.row_dimensions[row[0].row].height=130
    dv=DataValidation(type='list',formula1='"О,Д,Отп,Б,От,Вых"');p.add_data_validation(dv);dv.add(f'C5:{get_column_letter(p.max_column)}{p.max_row}')
    fact=Workbook();f=fact.active;f.title='СКУД_факт'
    f.append(['СКУД_факт — PPL Group (синтетические данные)']);f.append(['Период',period])
    f.append(['Формат: приход / уход / -- / длительность присутствия. Перерыв не вычитается.'])
    while f.max_row<8:f.append([''])
    f.append(['№','Ф И']+[d.day for d in dates]+['Итого за месяц'])
    totals=[0]*len(dates)
    for n,(name,codes) in enumerate(sorted(people),1):
        vals=[];total=0
        for j,code in enumerate(codes):
            v='—'
            if (code=='О' and rng.random()>.07) or (code!='О' and rng.random()<.015):
                start=540+rng.choice([-10,0,0,0,5,15,30]);end=1080+rng.choice([-30,0,0,15])
                if rng.random()<.025:v=f'{start//60:02}:{start%60:02}\n—\n--\n*'
                else:
                    minutes=end-start;total+=minutes;totals[j]+=minutes
                    v=f'{start//60:02}:{start%60:02}\n{end//60:02}:{end%60:02}\n--\n{minutes//60}:{minutes%60:02}'
            vals.append(v)
        dirty=('  '+name.upper().replace(' ','   ')+' ') if n%7==0 else name
        f.append([n,dirty]+vals+[total/1440]);f.cell(f.max_row,len(dates)+3).number_format='[h]:mm'
    f.append(['Итого','']+[v/1440 for v in totals]+[sum(totals)/1440]);style(f,9,True)
    f.column_dimensions['A'].width=8
    f.column_dimensions['B'].width=32
    for r in range(10,f.max_row+1):f.cell(r,len(dates)+3).number_format='[h]:mm'
    for c in range(3,len(dates)+3):f.cell(f.max_row,c).number_format='[h]:mm'
    return bytes_wb(plan),bytes_wb(fact)

