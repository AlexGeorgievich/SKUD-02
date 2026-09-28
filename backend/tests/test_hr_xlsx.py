import io
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend.app.config import Settings
from backend.app.hr.models import Base
from backend.app.hr.xlsx import HR_HEADERS, export_hr_xlsx, parse_hr_xlsx
from backend.app.main import create_app


def workbook_bytes(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(HR_HEADERS))
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class HrXlsxTests(unittest.TestCase):
    def make_client(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        app = create_app(Settings(data_dir=Path(temporary.name), database_url='sqlite+pysqlite:///:memory:'))
        Base.metadata.create_all(app.state.container.hr.repository.engine)
        self.addCleanup(app.state.container.hr.repository.engine.dispose)
        app.state.container.auth.create_user('hr', '123456789012', 'hr')
        client = TestClient(app, base_url='http://localhost')
        client.post('/api/login', json={'username': 'hr', 'password': '123456789012'})
        return client

    def test_exact_headers_dates_and_empty_values(self):
        values = [None] * 25
        values[0] = 'Иванова Анна Петровна'
        values[3] = 'Buying'
        values[7] = datetime(2026, 1, 12)
        values[11] = date(1990, 5, 21)
        parsed = parse_hr_xlsx(workbook_bytes([values]))
        self.assertEqual(parsed[0]['name'], 'Иванова Анна Петровна')
        self.assertEqual(parsed[0]['hire_date'], '2026-01-12')
        self.assertEqual(parsed[0]['birth_date'], '1990-05-21')
        self.assertIsNone(parsed[0]['telegram'])

    def test_export_keeps_original_25_column_structure(self):
        blob = export_hr_xlsx([{
            'plan_name': 'Иванова Анна Петровна', 'position': 'Специалист',
            'department': 'Buying', 'hire_date': date(2026, 1, 12),
            'birth_date': date(1990, 5, 21), 'telegram': '@anna_test',
        }])
        sheet = load_workbook(io.BytesIO(blob), read_only=True, data_only=True).active
        self.assertEqual(tuple(cell.value for cell in next(sheet.iter_rows(max_row=1)))[:25], HR_HEADERS)
        self.assertEqual(sheet.max_column, 25)
        self.assertEqual(parse_hr_xlsx(blob)[0]['telegram'], '@anna_test')

    def test_preview_reports_only_bad_contact_fields_and_apply_is_explicit(self):
        client = self.make_client()
        values = [None] * 25
        values[0], values[3], values[10], values[11] = 'Иванова Анна Петровна', 'Buying', 'май', '21.05'
        values[13], values[14], values[15] = 'bad telegram!', '12', 'bad-email'
        preview = client.post('/api/hr/xlsx/preview', files={'file': ('hr.xlsx', workbook_bytes([values]), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
        self.assertEqual(preview.status_code, 200, preview.text)
        body = preview.json()
        self.assertEqual(body['error_count'], 3)
        self.assertEqual({item['field'] for item in body['errors']}, {'telegram', 'personal_phone', 'work_email'})
        self.assertEqual(client.get('/api/hr/employees').json()['count'], 0)
        self.assertEqual(client.post('/api/hr/xlsx/apply', json={'batch_id': body['batch_id'], 'confirm_archive_ids': []}).status_code, 400)

    def test_full_file_requires_archive_confirmation_and_apply_is_idempotent(self):
        client = self.make_client()
        client.post('/api/hr/employees', json={'plan_name': 'Старова Анна', 'department': 'Legacy'})
        old_id = client.get('/api/hr/employees').json()['items'][0]['id']
        values = [None] * 25
        values[0], values[1], values[3] = 'Новая Елена Петровна', 'Специалист', 'Buying'
        values[7], values[9], values[13], values[14], values[15] = date(2026, 8, 1), 'ж', 't.me/new_test', '8 (999) 111-22-33', 'NEW@example.com'
        preview = client.post('/api/hr/xlsx/preview', files={'file': ('hr.xlsx', workbook_bytes([values]), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}).json()
        self.assertEqual(preview['error_count'], 0)
        self.assertEqual(preview['archive_candidates'], [old_id])
        self.assertEqual(client.post('/api/hr/xlsx/apply', json={'batch_id': preview['batch_id'], 'confirm_archive_ids': []}).status_code, 400)
        applied = client.post('/api/hr/xlsx/apply', json={'batch_id': preview['batch_id'], 'confirm_archive_ids': [old_id]})
        self.assertEqual(applied.status_code, 200, applied.text)
        self.assertEqual(applied.json(), {'created': 1, 'updated': 0, 'archived': 1, 'status': 'applied'})
        repeated = client.post('/api/hr/xlsx/apply', json={'batch_id': preview['batch_id'], 'confirm_archive_ids': [old_id]})
        self.assertEqual(repeated.json()['status'], 'already_applied')
        exported = client.get('/api/hr/xlsx/export')
        self.assertEqual(exported.status_code, 200)
        row = parse_hr_xlsx(exported.content)[0]
        self.assertEqual(row['name'], 'Новая Елена')
        employee = client.get('/api/hr/employees').json()['items'][0]
        self.assertIsNotNone(employee['position_id'])
        positions = client.get('/api/hr/catalogs').json()['positions']
        self.assertEqual([item['label'] for item in positions], ['Специалист'])
        self.assertEqual(row['telegram'], '@new_test')
        self.assertEqual(row['personal_phone'], '+79991112233')
        self.assertEqual(row['work_email'], 'new@example.com')


if __name__ == '__main__':
    unittest.main()
