import io
import json
import os
import tempfile
import unittest
from datetime import date
from pathlib import Path
from openpyxl import load_workbook
from backend.tests import support as core

class CoreTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.p,cls.f=core.generate('2026-08');cls.key=b'test-only-key'
    def test_population(self):
        self.assertEqual(sum(core.DEPTS.values()),166)
        self.assertEqual(len(core.read_input(self.p,'plan','2026-08')),166)
        self.assertEqual(len(core.read_input(self.f,'fact','2026-08')),166)
    def test_calendar(self):
        self.assertEqual(len(core.month_dates('2024-02')),29)
        self.assertEqual(len(core.month_dates('2025-02')),28)
    def test_normalization(self):
        self.assertEqual(core.token(' Иванов\u00a0  ИВАН ',self.key),core.token('иванов иван',self.key))
        self.assertNotEqual(core.token('Семёнов Иван',self.key),core.token('Семенов Иван',self.key))
    def test_period_reject(self):
        with self.assertRaises(ValueError):core.read_input(self.p,'plan','2026-07')
    def test_duplicate_reject(self):
        wb=load_workbook(io.BytesIO(self.p));wb.active['B6']=wb.active['B5'].value
        with self.assertRaises(ValueError):core.read_input(core.bytes_wb(wb),'plan','2026-08')
    def test_incomplete_and_invalid(self):
        self.assertIsNone(core.parse_event('09:00\n—\n--\n*')['minutes'])
        self.assertTrue(core.parse_event('25:00\n18:00\n--\n8:00')['issue'])
        self.assertEqual(core.parse_event('09:00\n18:00\n--\n9:00')['minutes'],540)
        self.assertTrue(core.parse_event('09:00\n18:00\n--\n8:00')['issue'])
    def test_pipeline(self):
        r,m=core.process(self.p,self.f,'2026-08',self.key,date(2026,8,31))
        self.assertEqual(len(r['days']),166*31);self.assertEqual(r['unmatched'],[])
        self.assertEqual(sum(e['minutes'] for e in r['employees']),sum(d['minutes'] or 0 for d in r['days']))
        self.assertTrue(any(d['problem'] for d in r['days']))
        for name in m.values():self.assertNotIn(name,json.dumps(r,ensure_ascii=False))
        wb=load_workbook(io.BytesIO(core.export(r,m)))
        self.assertEqual(wb.sheetnames,['План-факт','За месяц','Проверка данных','По дням'])
        self.assertEqual(wb['За месяц'].max_row,167)
        self.assertIn(wb['За месяц']['B2'].value,m.values())
        anon=load_workbook(io.BytesIO(core.export(r)))
        self.assertTrue(anon['За месяц']['B2'].value.startswith('EMP-'))
    def test_future_not_absence(self):
        r,_=core.process(self.p,self.f,'2026-08',self.key,date(2026,7,31))
        self.assertFalse(any('нет регистрации' in d['result'] for d in r['days']))
    def test_remote_not_absence(self):
        r,_=core.process(self.p,self.f,'2026-08',self.key,date(2026,8,31))
        self.assertFalse(any('нет регистрации' in d['result'] for d in r['days'] if d['plan']=='Д'))

if __name__=='__main__':unittest.main(verbosity=2)
