import io
import re
import tempfile
import unittest
from pathlib import Path
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from backend.app.config import Settings, PROJECT_ROOT
from backend.app.main import create_app

class ApiTests(unittest.TestCase):
    def test_full_workflow_and_roles(self):
        with tempfile.TemporaryDirectory() as tmp:
            app = create_app(Settings(data_dir=Path(tmp)))
            container = app.state.container
            container.auth.create_user('admin', 'testing-password-123', 'admin')
            client = TestClient(app, base_url='http://localhost')
            self.assertEqual(client.get('/api/result').status_code, 401)
            self.assertEqual(client.get('/api/health').json()['version'], '0.3.0')
            html = client.get('/')
            self.assertEqual(html.status_code, 200)
            self.assertIn('id="root"', html.text)
            asset = re.search(r'src="([^"]+\.js)"', html.text).group(1)
            self.assertEqual(client.get(asset).status_code, 200)
            self.assertEqual(client.get('/static/react/../../data/secret.key').status_code, 404)
            self.assertEqual(client.post('/api/login', json={'username': [], 'password': 'x'}).status_code, 422)
            self.assertEqual(client.post('/api/login', json={'username': 'admin', 'password': 'wrong'}).status_code, 401)
            self.assertEqual(client.post('/api/login', json={'username': 'admin', 'password': 'testing-password-123'}).status_code, 200)
            payload = {'period': '2026-08', 'asof': '2026-08-31'}
            blobs = {'plan': ('plan.xlsx', (PROJECT_ROOT/'demo/plan.xlsx').read_bytes()), 'fact': ('fact.xlsx', (PROJECT_ROOT/'demo/skud_fact.xlsx').read_bytes())}
            preview = client.post('/api/import', data={**payload, 'preview': 'true'}, files=blobs)
            self.assertEqual(preview.status_code, 200, preview.text)
            self.assertEqual(preview.json()['plan_count'], 166)
            response = client.post('/api/import', data=payload, files=blobs)
            self.assertEqual(response.status_code, 200, response.text)
            result = client.get('/api/result').json()
            self.assertEqual(len(result['employees']), 166)
            first = result['employees'][0]
            exported = client.get('/api/export', params={'search': first['name']})
            wb = load_workbook(io.BytesIO(exported.content))
            self.assertEqual(wb['За месяц'].max_row, 2)
            self.assertEqual(wb['За месяц']['B2'].value, first['name'])
            snapshot, identities = container.repository.load_snapshot()
            self.assertEqual(len(identities), 166)
            self.assertNotIn('name', snapshot['employees'][0])
            # Separate instances must not share sessions or mutate global paths.
            second = create_app(Settings(data_dir=Path(tmp)/'other'))
            with self.assertRaises(Exception):
                second.state.container.auth.current_user(client.cookies.get('tt_session'))
            container.auth.create_user('head', 'testing-password-123', 'manager', 'HR')
            container.auth.create_user('employee', 'testing-password-123', 'employee', employee=first['name'])
            container.auth.create_user('audit', 'testing-password-123', 'auditor')
            for login, count in [('head', 4), ('employee', 1), ('audit', 166)]:
                client.post('/api/logout')
                self.assertEqual(client.get('/api/result').status_code, 401)
                self.assertEqual(client.post('/api/login', json={'username': login, 'password': 'testing-password-123'}).status_code, 200)
                scoped = client.get('/api/result').json()
                self.assertEqual(len(scoped['employees']), count)
                self.assertEqual(client.post('/api/import', data=payload, files=blobs).status_code, 403)
                if login == 'head':
                    wb = load_workbook(io.BytesIO(client.get('/api/export?department=Buying').content))
                    self.assertEqual(wb['За месяц'].max_row, 1)
                if login == 'audit':
                    self.assertTrue(all(e['name'].startswith('EMP-') for e in scoped['employees']))
                    self.assertEqual(client.get('/api/audit').status_code, 200)
